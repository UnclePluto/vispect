from io import BytesIO
import pytest
from openpyxl import load_workbook
from apps.core.context import tenant_transaction
from apps.core.errors import ApiError

def workbook(data):return load_workbook(BytesIO(data))
def test_doctor_export_excludes_other_patients(world):
    from apps.followups.exports import export_answers
    from apps.followups.answers import submit_answer
    from apps.followups.enrollments import add_subject,confirm_enrollment
    from apps.studies.services import set_study_members
    a=world.actor('a','doctor');admin=world.actor('a','tenant_admin');e,v,b=world.entry('a');submit_answer(a,v.id,b.id,{'q1':'我的答案'},'s')
    other=world.doctor('a');p=world.patient('a',doctor_ids=[other.id])
    with tenant_transaction(a.tenant_id):s=e.study
    set_study_members(admin,s.id,[other.id],s.revision)
    other_e=add_subject(admin,s.id,p.id,other.id,'other','other-add');other_e=confirm_enrollment(admin,other_e.id,'other-code',other_e.added_on,'other-confirm')
    with tenant_transaction(a.tenant_id):other_visit=other_e.current_visit;other_binding=other_visit.stage.bindings.first()
    submit_answer(admin,other_visit.id,other_binding.id,{'q1':'其他患者答案'},'other-answer')
    own=workbook(export_answers(a,study_id=s.id));all_rows=list(own.active.values)
    assert own.active.max_row==2 and '其他患者答案' not in str(all_rows)
    assert workbook(export_answers(admin,study_id=s.id)).active.max_row==3

def test_formula_text_and_leading_zero_are_preserved(world):
    from apps.followups.exports import export_answers
    from apps.followups.answers import submit_answer
    e,v,b=world.entry('a');a=world.actor('a','doctor')
    with tenant_transaction(a.tenant_id):patient=e.patient
    patient.code='000001';world.owner_save(patient)
    answer=submit_answer(a,v.id,b.id,{'q1':'=SUM(1,2)'},'s');sheet=workbook(export_answers(a,answer_id=answer.id)).active
    assert sheet.cell(2,1).value=='000001' and sheet.cell(2,1).data_type=='s'
    assert sheet.cell(2,7).value=='=SUM(1,2)' and sheet.cell(2,7).data_type=='s'

def test_empty_study_and_multiple_versions_have_legal_sheets(world):
    from apps.followups.exports import export_answers
    from apps.followups.answers import submit_answer
    from apps.followups.enrollments import add_subject,confirm_enrollment
    from apps.studies.services import create_study,save_protocol,set_study_state
    a=world.actor('a','doctor');s=world.study('a');assert workbook(export_answers(a,study_id=s.id)).active.max_row==1
    definition={'protocol_version':1,'title':'长标题'*100,'questions':[{'id':'q','type':'number','title':'数值','required':False}]}
    v1=world.form_version('a',definition);v2=world.form_version('a',definition)
    s=create_study(a,{'name':'多版本','code':'versions'});s=save_protocol(a,s.id,[{'kind':'baseline','name':'基线','offset_days':0,'bindings':[{'form_version_id':str(v1.id),'respondent':'doctor','required':True},{'form_version_id':str(v2.id),'respondent':'patient','required':True}]}],s.revision);s=set_study_state(a,s.id,'active',s.revision)
    p=world.patient('a');e=add_subject(a,s.id,p.id,a.doctor_id,'obs','add');e=confirm_enrollment(a,e.id,'000001',e.added_on,'confirm')
    with tenant_transaction(a.tenant_id):visit=e.current_visit;bindings=list(visit.stage.bindings.all())
    submit_answer(a,visit.id,bindings[0].id,{'q':0},'s1');submit_answer(a,visit.id,bindings[1].id,{},'s2')
    book=workbook(export_answers(a,study_id=s.id));assert len(set(book.sheetnames))==2 and all(len(n)<=31 for n in book.sheetnames)
    assert sorted([str(sheet.cell(2,7).value) for sheet in book.worksheets])==['0','None']

def test_revoked_access_rejects_old_download_and_revision_requests(world):
    from apps.followups.answers import submit_answer,revise_answer
    from apps.followups.enrollments import transfer_followup
    from apps.studies.services import set_study_members
    from apps.patients.services import assign_doctors
    e,v,b=world.entry('a');a=world.actor('a','doctor');admin=world.actor('a','tenant_admin');answer=submit_answer(a,v.id,b.id,{'q1':'原'},'s');answer=revise_answer(a,answer.id,{'q1':'新'},'纠正',answer.revision);client=world.login('a','doctor')
    assert client.get(f'/api/answers/{answer.id}/export/').status_code==200
    other=world.doctor('a')
    with tenant_transaction(a.tenant_id):s=e.study;p=e.patient
    set_study_members(admin,s.id,[other.id],s.revision);p=assign_doctors(admin,p.id,[a.doctor_id,other.id],p.revision);transfer_followup(admin,e.id,other.id,e.revision);assign_doctors(admin,p.id,[other.id],p.revision)
    assert client.get(f'/api/answers/{answer.id}/export/').status_code==404
    assert client.get(f'/api/answers/{answer.id}/revisions/').status_code==404

def test_long_text_export_preserves_full_content_and_audit(world):
    from apps.followups.exports import export_answers
    from apps.followups.answers import submit_answer
    from apps.core.models import AuditEvent
    e,v,b=world.entry('a');a=world.actor('a','doctor');text='😀'*20000+'完整结尾'
    answer=submit_answer(a,v.id,b.id,{'q1':text},'s');sheet=workbook(export_answers(a,answer_id=answer.id)).active
    assert sheet.cell(2,7).value+sheet.cell(2,8).value==text
    assert sheet.cell(2,7).data_type=='s' and sheet.cell(2,8).data_type=='s'
    with tenant_transaction(a.tenant_id):event=AuditEvent.objects.filter(action='export',object_id=answer.id).get()
    assert event.changes['count']==1 and 'answers' not in event.changes

def test_export_does_not_expose_patient_names_in_filename(world):
    from apps.followups.answers import submit_answer
    e,v,b=world.entry('a');answer=submit_answer(world.actor('a','doctor'),v.id,b.id,{'q1':'答案'},'s')
    response=world.login('a','doctor').get(f'/api/answers/{answer.id}/export/')
    assert response['Content-Disposition']=='attachment; filename="answer.xlsx"'
    assert response['Cache-Control']=='no-store'
    assert world.login(None,'platform_admin').get(f'/api/answers/{answer.id}/export/').status_code==403

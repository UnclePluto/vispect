import re
from collections import defaultdict
from io import BytesIO
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from apps.core.permissions import require_tenant
from apps.core.context import tenant_transaction
from apps.core.helpers import get_object
from apps.core.errors import ApiError
from apps.core.audit import audit
from apps.studies.selectors import visible_studies
from .selectors import visible_answers
from .serializers import recorder_name
HEADERS=['患者档案编号','研究入组编号','访视','填写对象','录入人','提交时间']
XML_INVALID=re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')

def safe_text(value):return XML_INVALID.sub(lambda m:f'\\u{ord(m.group()):04x}',value)
def text_parts(value):
    value=safe_text(value);parts=[];current=[];length=0
    for char in value:
        units=2 if ord(char)>0xffff else 1
        if length+units>32767:parts.append(''.join(current));current=[];length=0
        current.append(char);length+=units
    parts.append(''.join(current));return parts

def display_answer(question,answers):
    value=answers.get(question['id'])
    if value is None:return None
    if question['type'] in ['radio','select']:return next((o['label'] for o in question['options'] if o['id']==value),'')
    if question['type']=='checkbox':return '；'.join(o['label'] for o in question['options'] if o['id'] in value)
    return value

def write_row(sheet,row,values,header=False):
    for column,value in enumerate(values,1):
        cell=sheet.cell(row,column)
        if isinstance(value,str):cell.value=safe_text(value);cell.data_type='s'
        else:cell.value=value
        cell.alignment=Alignment(vertical='top',wrap_text=True)
        if header:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1677FF')

def sheet_name(book,version):
    stem=f'v{version.number}-{str(version.id)[:20]}'[:31];name=stem;index=2
    while name in book.sheetnames:name=f'{stem[:27]}-{index}';index+=1
    return name

def export_answers(actor,study_id=None,answer_id=None):
    require_tenant(actor)
    if (study_id is None)==(answer_id is None):raise ApiError('invalid_export_scope','导出须指定单份答案或一个研究',status=400)
    with tenant_transaction(actor.tenant_id):
        query=visible_answers(actor).select_related('form_version','recorded_by','recorded_by_patient','submitted_by','submitted_by_patient','visit__stage','visit__enrollment__patient','visit__enrollment__study')
        if study_id is not None:
            get_object(visible_studies(actor),study_id);query=query.filter(visit__study_id=study_id);target=study_id;object_type='study'
        else:
            get_object(visible_answers(actor),answer_id);query=query.filter(id=answer_id);target=answer_id;object_type='answer'
        rows=list(query.order_by('created_at','id'));groups=defaultdict(list)
        for answer in rows:groups[answer.form_version_id].append(answer)
        book=Workbook();book.remove(book.active)
        if not groups:sheet=book.create_sheet('答案');write_row(sheet,1,HEADERS+['答案状态','提交来源','原始提交人','最近修改时间'],True)
        for answers in groups.values():
            version=answers[0].form_version;questions=version.definition['questions'];sheet=book.create_sheet(sheet_name(book,version))
            values=[[display_answer(q,a.answers) for q in questions] for a in answers]
            widths=[max([len(text_parts(row[i])) if isinstance(row[i],str) else 1 for row in values],default=1) for i in range(len(questions))]
            headers=list(HEADERS)
            for q,count in zip(questions,widths):
                headers.extend([f"{q['title']} [{q['id']}]"+(f'（续 {part+1}）' if part else '') for part in range(count)])
            headers.extend(['答案状态','提交来源','原始提交人','最近修改时间']);write_row(sheet,1,headers,True)
            for index,(answer,question_values) in enumerate(zip(answers,values),2):
                e=answer.visit.enrollment
                data=[e.patient.code,e.enrollment_code,answer.visit.stage.name,('患者（工作人员代录）' if answer.is_proxy else '患者自填') if answer.respondent=='patient' else '医生',recorder_name(answer.recorded_by,answer.recorded_by_patient),timezone.localtime(answer.submitted_at).isoformat() if answer.submitted_at else None]
                for value,count in zip(question_values,widths):
                    parts=text_parts(value) if isinstance(value,str) else [value];data.extend(parts+[None]*(count-len(parts)))
                data.extend(['已提交' if answer.status=='submitted' else '草稿',{'patient':'患者自填','proxy':'工作人员代录','doctor':'医生填写'}.get(answer.submission_source,'尚未提交'),recorder_name(answer.submitted_by,answer.submitted_by_patient),timezone.localtime(answer.updated_at).isoformat()]);write_row(sheet,index,data)
        for sheet in book.worksheets:
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for column in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(column)].width=24 if column>6 else 22
        output=BytesIO();book.save(output)
        audit(actor,'export',object_type,target,{'count':len(rows),'format':'xlsx','version_count':len(groups)})
        return output.getvalue()
